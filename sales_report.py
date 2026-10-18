import csv
import json
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

BASE_DIR = Path(__file__).parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"
INPUT_FILE = INPUT_DIR / "sales.csv"

OUTPUT_DIR.mkdir(exist_ok=True)
INPUT_DIR.mkdir(exist_ok=True)

# Create sample data automatically if no input file exists.
if not INPUT_FILE.exists():
    sample_data = [
        ["date", "product", "quantity", "unit_price"],
        ["2026-01-05", "Keyboard", 5, 25.00],
        ["2026-01-12", "Mouse", 8, 12.50],
        ["2026-01-20", "Monitor", 2, 180.00],
        ["2026-02-03", "Keyboard", 7, 25.00],
        ["2026-02-14", "Headphones", 4, 45.00],
        ["2026-02-21", "Mouse", 10, 12.50],
        ["2026-03-02", "Monitor", 3, 180.00],
        ["2026-03-10", "Headphones", 6, 45.00],
    ]

    with INPUT_FILE.open("w", newline="", encoding="utf-8") as file:
        csv.writer(file).writerows(sample_data)

# Read and validate sales data.
sales = []
with INPUT_FILE.open("r", newline="", encoding="utf-8-sig") as file:
    reader = csv.DictReader(file)
    required = {"date", "product", "quantity", "unit_price"}

    if not reader.fieldnames or not required.issubset(reader.fieldnames):
        raise ValueError(
            "CSV must contain: date, product, quantity, unit_price"
        )

    for line_number, row in enumerate(reader, start=2):
        try:
            date = row["date"].strip()
            product = row["product"].strip()
            quantity = int(row["quantity"])
            unit_price = float(row["unit_price"])

            if not product or quantity < 0 or unit_price < 0:
                raise ValueError("Invalid product, quantity, or price")

            month = date[:7]
            if len(month) != 7 or date[4] != "-":
                raise ValueError("Date must use YYYY-MM-DD format")

            revenue = round(quantity * unit_price, 2)
            sales.append({
                "date": date,
                "month": month,
                "product": product,
                "quantity": quantity,
                "unit_price": unit_price,
                "revenue": revenue,
            })
        except (ValueError, TypeError, KeyError) as error:
            raise ValueError(
                f"Invalid data on CSV row {line_number}: {error}"
            ) from error

if not sales:
    raise ValueError("No valid sales records found in input/sales.csv")

# Calculate summaries.
product_totals = defaultdict(lambda: {"quantity": 0, "revenue": 0.0})
monthly_totals = defaultdict(float)

for sale in sales:
    product_totals[sale["product"]]["quantity"] += sale["quantity"]
    product_totals[sale["product"]]["revenue"] += sale["revenue"]
    monthly_totals[sale["month"]] += sale["revenue"]

total_revenue = round(sum(s["revenue"] for s in sales), 2)
total_units = sum(s["quantity"] for s in sales)
best_product = max(
    product_totals.items(),
    key=lambda item: item[1]["revenue"]
)

# Create formatted Excel workbook.
wb = Workbook()
summary = wb.active
summary.title = "Dashboard"

summary.append(["SALES AUTOMATION REPORT"])
summary.append(["Metric", "Value"])
summary.append(["Total Revenue", total_revenue])
summary.append(["Total Units Sold", total_units])
summary.append(["Sales Records", len(sales)])
summary.append(["Top Product by Revenue", best_product[0]])

products_ws = wb.create_sheet("Product Summary")
products_ws.append(["Product", "Units Sold", "Revenue"])

for product, values in sorted(product_totals.items()):
    products_ws.append([
        product,
        values["quantity"],
        round(values["revenue"], 2),
    ])

monthly_ws = wb.create_sheet("Monthly Summary")
monthly_ws.append(["Month", "Revenue"])

for month, revenue in sorted(monthly_totals.items()):
    monthly_ws.append([month, round(revenue, 2)])

details = wb.create_sheet("Sales Details")
details.append([
    "Date", "Month", "Product", "Quantity", "Unit Price", "Revenue"
])

for sale in sales:
    details.append([
        sale["date"], sale["month"], sale["product"],
        sale["quantity"], sale["unit_price"], sale["revenue"]
    ])

# Style worksheets.
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    for column_cells in ws.columns:
        width = max(
            len(str(cell.value or "")) for cell in column_cells
        )
        ws.column_dimensions[column_cells[0].column_letter].width = min(
            width + 3, 35
        )

for row in range(3, summary.max_row + 1):
    if summary.cell(row, 1).value == "Total Revenue":
        summary.cell(row, 2).number_format = "#,##0.00"

for ws in [products_ws, monthly_ws, details]:
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in (2, 5, 6) and isinstance(cell.value, (int, float)):
                if ws == products_ws and cell.column == 2:
                    continue
                cell.number_format = "#,##0.00"

# Add monthly revenue chart.
chart = BarChart()
chart.type = "col"
chart.title = "Monthly Revenue"
chart.y_axis.title = "Revenue"
chart.x_axis.title = "Month"
chart.add_data(
    Reference(monthly_ws, min_col=2, min_row=1, max_row=monthly_ws.max_row),
    titles_from_data=True,
)
chart.set_categories(
    Reference(monthly_ws, min_col=1, min_row=2, max_row=monthly_ws.max_row)
)
chart.height = 8
chart.width = 15
summary.add_chart(chart, "D2")

excel_path = OUTPUT_DIR / "sales_report.xlsx"
wb.save(excel_path)

json_path = OUTPUT_DIR / "summary_report.json"
report = {
    "records_processed": len(sales),
    "total_revenue": total_revenue,
    "total_units_sold": total_units,
    "top_product_by_revenue": best_product[0],
    "monthly_revenue": {
        month: round(revenue, 2)
        for month, revenue in sorted(monthly_totals.items())
    },
}
json_path.write_text(json.dumps(report, indent=2), encoding="utf-8")

print("\nSUCCESS: Sales report generated!")
print(f"Records processed: {len(sales)}")
print(f"Total revenue: {total_revenue:.2f}")
print(f"Total units sold: {total_units}")
print(f"Top product: {best_product[0]}")
print(f"\nExcel report: {excel_path}")
print(f"JSON summary: {json_path}")